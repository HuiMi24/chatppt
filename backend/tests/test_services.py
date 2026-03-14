import tempfile
import unittest
from pathlib import Path

from pptx import Presentation

from app.chat_service import ChatPlanner
from app.models import EditInstruction
from app.ppt_service import PPTService


class TestPPTServices(unittest.TestCase):
    def setUp(self):
        self.tmpdir = tempfile.TemporaryDirectory()
        self.base = Path(self.tmpdir.name) / "sample.pptx"

        prs = Presentation()
        title_slide = prs.slides.add_slide(prs.slide_layouts[0])
        title_slide.shapes.title.text = "Old Title"
        title_slide.placeholders[1].text = "Old Subtitle"
        prs.save(self.base)

        self.ppt = PPTService()
        self.chat = ChatPlanner(self.ppt)

    def tearDown(self):
        self.tmpdir.cleanup()

    def test_parse_and_edit(self):
        doc = self.ppt.parse_ppt(str(self.base))
        self.assertEqual(doc.slide_count, 1)
        self.assertTrue(len(doc.slides[0].shapes) >= 1)

        title_shape = next(s for s in doc.slides[0].shapes if s.role == "title")
        out = self.ppt.apply_edits(
            str(self.base),
            [
                EditInstruction(
                    slide_index=0,
                    shape_index=title_shape.shape_index,
                    new_text="New Title",
                )
            ],
        )
        out_doc = self.ppt.parse_ppt(out)
        self.assertEqual(out_doc.slides[0].title, "New Title")

    def test_rule_based_chat(self):
        plan, used_llm = self.chat.build_plan(str(self.base), "把第1页标题改为年度复盘")
        self.assertFalse(used_llm)
        self.assertEqual(len(plan), 1)


if __name__ == "__main__":
    unittest.main()
